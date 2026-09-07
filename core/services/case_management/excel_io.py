"""用例管理 Excel 导入导出（独立实现，不改动工作台导出）。"""

from __future__ import annotations

import io
import re
from typing import Any, BinaryIO

from openpyxl import Workbook, load_workbook

from core.services.case_management.case_db import bulk_insert_cases, list_cases
from core.services.case_management.suite_db import assert_suite_in_project, list_suites


HEADERS = [
    "用例名称",
    "所属模块",
    "优先级",
    "状态",
    "前置条件",
    "步骤描述",
    "预期结果",
    "标签",
]


def _split_numbered(text: str) -> list[str]:
    raw = str(text or "").replace("\\n", "\n").strip()
    if not raw:
        return []
    parts = re.split(r"(?:^|\n)\s*(?:\[\d+\]|\d+[\.、．)])\s*", raw)
    out = [p.strip() for p in parts if p and p.strip()]
    if out:
        return out
    return [raw]


def _join_numbered(items: list[str]) -> str:
    lines = []
    for i, item in enumerate(items, 1):
        lines.append("[%d] %s" % (i, item))
    return "\n".join(lines)


def export_project_excel(
    user_id: str, project_id: str, *, suite_id: str | None = None
) -> bytes:
    from core.services.case_management.access import assert_project_owner
    from core.services.case_management.suite_schema_db import get_schema_by_suite

    assert_project_owner(user_id, project_id)
    suites = {s["id"]: s for s in list_suites(user_id, project_id)}
    sid_filter = str(suite_id or "").strip()
    if sid_filter and sid_filter not in ("__all__", "*"):
        assert_suite_in_project(user_id, project_id, sid_filter)
        list_suite_id: str = sid_filter
    else:
        list_suite_id = "__all__"

    # 拉全量（分页合并）
    page = 1
    all_cases: list[dict[str, Any]] = []
    while True:
        batch = list_cases(
            user_id,
            project_id,
            suite_id=list_suite_id,
            page=page,
            page_size=100,
            full=True,
        )
        items = batch.get("items") or []
        all_cases.extend(items)
        if len(all_cases) >= int(batch.get("total") or 0) or not items:
            break
        page += 1
        if page > 200:
            break

    schema_cols = None
    if list_suite_id != "__all__":
        sch = get_schema_by_suite(list_suite_id)
        if sch and str(sch.get("status") or "") == "locked" and sch.get("columns"):
            schema_cols = sch.get("columns") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "用例"
    if schema_cols:
        headers = [str(c.get("label") or c.get("key") or "") for c in schema_cols]
        ws.append(headers)
        for case in all_cases:
            fields = case.get("fields") or {}
            if not isinstance(fields, dict):
                fields = {}
            row_out = []
            for col in schema_cols:
                key = str(col.get("key") or "")
                role = str(col.get("role") or "")
                if key in fields and fields.get(key) not in (None, ""):
                    row_out.append(str(fields.get(key) or ""))
                elif role == "title":
                    row_out.append(str(case.get("title") or ""))
                elif role == "priority":
                    row_out.append(str(case.get("priority") or "P2"))
                elif role == "status":
                    row_out.append(str(case.get("status") or "draft"))
                elif role == "precondition":
                    row_out.append(str(case.get("precondition") or ""))
                elif role == "tags":
                    row_out.append(",".join(str(t) for t in (case.get("tags") or [])))
                elif role == "steps":
                    steps = case.get("steps") or []
                    row_out.append(_join_numbered([str(s.get("step") or "") for s in steps]))
                elif role == "expect":
                    steps = case.get("steps") or []
                    row_out.append(
                        _join_numbered([str(s.get("expect") or "") for s in steps])
                    )
                else:
                    row_out.append(str(fields.get(key) or ""))
            ws.append(row_out)
    else:
        ws.append(HEADERS)
        for case in all_cases:
            suite_name = ""
            sid = case.get("suite_id")
            if sid and sid in suites:
                suite_name = str(suites[sid].get("name") or "")
            steps = case.get("steps") or []
            step_texts = [str(s.get("step") or "") for s in steps]
            expect_texts = [str(s.get("expect") or "") for s in steps]
            tags = case.get("tags") or []
            ws.append(
                [
                    case.get("title") or "",
                    suite_name,
                    case.get("priority") or "P2",
                    case.get("status") or "draft",
                    case.get("precondition") or "",
                    _join_numbered(step_texts),
                    _join_numbered(expect_texts),
                    ",".join(str(t) for t in tags),
                ]
            )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_project_excel(
    user_id: str,
    project_id: str,
    file_obj: BinaryIO,
    *,
    suite_id: str,
) -> dict[str, Any]:
    """将 Excel 用例全部导入到指定目录（suite_id 必填，不再按模块自动建目录）。"""
    target = assert_suite_in_project(user_id, project_id, suite_id)
    target_suite_id = str(target["id"])

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 为空")
    header = [str(c or "").strip() for c in rows[0]]
    col = {name: idx for idx, name in enumerate(header)}
    for required in ("用例名称",):
        if required not in col:
            raise ValueError("缺少必要列：用例名称")

    prepared: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if not raw:
            continue

        def cell(name: str) -> str:
            idx = col.get(name)
            if idx is None or idx >= len(raw):
                return ""
            return str(raw[idx] or "").strip()

        title = cell("用例名称")
        if not title:
            continue
        module = cell("所属模块")
        tags = [t.strip() for t in cell("标签").split(",") if t.strip()]
        if module and module not in tags:
            tags.append(module)

        step_parts = _split_numbered(cell("步骤描述"))
        expect_parts = _split_numbered(cell("预期结果"))
        n = max(len(step_parts), len(expect_parts), 1 if title else 0)
        steps = []
        for i in range(n):
            steps.append(
                {
                    "step": step_parts[i] if i < len(step_parts) else "",
                    "expect": expect_parts[i] if i < len(expect_parts) else "",
                }
            )
        priority = (cell("优先级") or "P2").upper()
        if priority not in ("P0", "P1", "P2", "P3"):
            priority = "P2"
        status = (cell("状态") or "draft").lower()
        if status not in ("draft", "ready", "deprecated"):
            status = "draft"
        prepared.append(
            {
                "title": title,
                "suite_id": target_suite_id,
                "priority": priority,
                "status": status,
                "precondition": cell("前置条件"),
                "steps": steps,
                "tags": tags,
                "source": "excel",
            }
        )

    result = bulk_insert_cases(user_id, project_id, prepared)
    result["parsed"] = len(prepared)
    result["suite_id"] = target_suite_id
    result["suite_name"] = target.get("name") or ""
    return result