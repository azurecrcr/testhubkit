"""带目录列结构的导入编排（新入口，不改动旧 excel_io / workbench_import 默认行为）。"""

from __future__ import annotations

import io
import json
from typing import Any, BinaryIO, Optional

from openpyxl import load_workbook

from core.services.case_management.case_db import (
    bulk_insert_cases,
    get_case,
    update_case,
    _normalize_fields,
    _normalize_steps,
    _normalize_tags,
)
from core.services.case_management.case_ops_db import titles_in_suite
from core.services.case_management.suite_db import assert_suite_in_project
from core.services.case_management import suite_schema_db as schema_db
from core.services.test_cases.requirement_case_db import get_requirement_case

_ALLOWED_DUP_MODES = frozenset({"create", "skip", "upsert"})


def _import_case_compare_snapshot(item: dict[str, Any] | None) -> dict[str, Any]:
    """导入比对用快照：只含覆盖更新会写入的业务字段。"""
    item = item or {}
    priority = str(item.get("priority") or "P2").strip().upper() or "P2"
    if priority not in ("P0", "P1", "P2", "P3"):
        priority = "P2"
    status = str(item.get("status") or "draft").strip().lower() or "draft"
    if status not in ("draft", "ready", "deprecated"):
        status = "draft"
    return {
        "title": str(item.get("title") or "").strip()[:500],
        "priority": priority,
        "status": status,
        "precondition": str(item.get("precondition") or ""),
        "steps": _normalize_steps(item.get("steps")),
        "tags": _normalize_tags(item.get("tags")),
        "fields": _normalize_fields(item.get("fields")),
    }


def _upsert_prepared_case_counted(
    user_id: str,
    case_id: str,
    item: dict[str, Any],
) -> str:
    """
    覆盖更新单条并返回真实结果：
    - updated: 内容有变化且写入成功
    - unchanged: 同名且内容完全一致（不算更新成功）
    - failed: 用例不存在/已删或写入失败
    不改动 update_case 本体，避免影响其它编辑入口。
    """
    existing = get_case(case_id)
    if not existing or existing.get("is_deleted"):
        return "failed"
    payload = {
        "title": str((item or {}).get("title") or "").strip(),
        "priority": (item or {}).get("priority") or "P2",
        "status": (item or {}).get("status") or "draft",
        "precondition": (item or {}).get("precondition") or "",
        "steps": (item or {}).get("steps") or [],
        "tags": (item or {}).get("tags") or [],
        "fields": (item or {}).get("fields") or {},
    }
    if _import_case_compare_snapshot(payload) == _import_case_compare_snapshot(existing):
        return "unchanged"
    try:
        update_case(user_id, case_id, payload)
    except Exception:
        return "failed"
    after = get_case(case_id)
    if not after or after.get("is_deleted"):
        return "failed"
    if _import_case_compare_snapshot(payload) != _import_case_compare_snapshot(after):
        return "failed"
    return "updated"


def _apply_prepared_cases(
    user_id: str,
    project_id: str,
    suite_id: str,
    prepared: list[dict[str, Any]],
    *,
    duplicate_mode: str = "create",
) -> dict[str, Any]:
    """按 duplicate_mode 写入；默认 create 走既有 bulk_insert，行为不变。"""
    mode = str(duplicate_mode or "create").strip().lower() or "create"
    if mode not in _ALLOWED_DUP_MODES:
        raise ValueError("导入模式仅支持 create/skip/upsert")
    if mode == "create":
        result = bulk_insert_cases(user_id, project_id, prepared)
        result.setdefault("skipped", 0)
        result.setdefault("updated", 0)
        result.setdefault("unchanged", 0)
        result.setdefault("failed", 0)
        result["duplicate_mode"] = mode
        return result

    existing = titles_in_suite(user_id, project_id, suite_id)
    to_create: list[dict[str, Any]] = []
    updated = 0
    skipped = 0
    unchanged = 0
    failed = 0
    errors: list[str] = []
    for i, item in enumerate(prepared or []):
        title = str((item or {}).get("title") or "").strip()
        if not title:
            continue
        old_id = existing.get(title)
        if not old_id:
            to_create.append(item)
            continue
        if mode == "skip":
            skipped += 1
            continue
        status = _upsert_prepared_case_counted(user_id, old_id, item)
        if status == "updated":
            updated += 1
        elif status == "unchanged":
            unchanged += 1
        else:
            failed += 1
            errors.append("第 %d 条更新失败: %s" % (i + 1, title))
            if len(errors) >= 20:
                break
    created_part = (
        bulk_insert_cases(user_id, project_id, to_create)
        if to_create
        else {"created": 0, "errors": []}
    )
    created_errors = created_part.get("errors") or []
    errors.extend(created_errors)
    created_n = int(created_part.get("created") or 0)
    create_failed = max(0, len(to_create) - created_n)
    failed += create_failed
    return {
        "created": created_n,
        "updated": updated,
        "skipped": skipped,
        "unchanged": unchanged,
        "failed": failed,
        "errors": errors[:20],
        "duplicate_mode": mode,
    }


def _read_excel_table(file_obj: BinaryIO) -> tuple[list[str], list[list[Any]]]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 为空")
    headers = [str(c or "").strip() for c in rows[0]]
    if not any(headers):
        raise ValueError("Excel 表头为空")
    body: list[list[Any]] = []
    for raw in rows[1:]:
        if not raw:
            continue
        body.append(list(raw))
    return headers, body


def _load_workbench_table(
    user_id: str, *, lanhu_pid: str, lanhu_doc_id: str, lanhu_page_id: str
) -> tuple[list[str], list[list[Any]], dict[str, Any]]:
    src = get_requirement_case(
        user_id,
        lanhu_pid=str(lanhu_pid or ""),
        lanhu_doc_id=str(lanhu_doc_id or ""),
        lanhu_page_id=str(lanhu_page_id or ""),
    )
    if not src:
        raise ValueError("未找到工作台用例数据")
    payload = src.get("payload") or {}
    if isinstance(payload, str):
        try:
            payload = json.loads(payload)
        except json.JSONDecodeError:
            payload = {}
    columns = [str(c or "").strip() for c in (payload.get("columns") or [])]
    rows_raw = payload.get("rows") or []
    if not columns:
        raise ValueError("工作台该页无列头")
    if not isinstance(rows_raw, list) or not rows_raw:
        raise ValueError("工作台该页无用例行")
    rows: list[list[Any]] = []
    for row in rows_raw:
        if isinstance(row, list):
            rows.append(row)
    return columns, rows, src


def preview_from_headers(
    user_id: str,
    project_id: str,
    suite_id: str,
    headers: list[str],
    *,
    sample_rows: Optional[list[list[Any]]] = None,
) -> dict[str, Any]:
    assert_suite_in_project(user_id, project_id, suite_id)
    # 空目录自动继承父级/兄弟已锁定表头（如 Web/1 锁定后，Web/2 沿用）
    schema = schema_db.ensure_inherited_schema(user_id, suite_id)
    headers_clean = [str(h).strip() for h in headers if str(h).strip()]
    sample: list[dict[str, str]] = []
    for row in (sample_rows or [])[:3]:
        sample.append(schema_db.row_values_by_header(headers_clean, list(row)))

    if str(schema.get("status") or "") != "locked" or not (schema.get("columns") or []):
        proposed = schema_db.build_columns_from_headers(headers_clean)
        # build_columns_from_headers 内已校验上限；此处再兜底一次
        schema_db.assert_schema_column_limit(proposed)
        return {
            "mode": "define_schema",
            "schema": schema,
            "proposed_columns": proposed,
            "source_headers": headers_clean,
            "sample_rows": sample,
            "suggested_mapping": {},
            "missing_required_keys": [],
            "unmapped_headers": [],
            "max_columns": schema_db.MAX_SCHEMA_COLUMNS,
        }

    # 已锁定：严格固定表头（不一致直接拦截，不再走字段映射）
    strict = schema_db.compare_locked_headers_strict(
        headers_clean, schema.get("columns") or []
    )
    return {
        "mode": strict.get("mode") or "header_mismatch",
        "schema": schema,
        "proposed_columns": schema.get("columns") or [],
        "source_headers": headers_clean,
        "expected_headers": strict.get("expected_headers") or [],
        "sample_rows": sample,
        "suggested_mapping": strict.get("suggested_mapping") or {},
        "missing_headers": strict.get("missing_headers") or [],
        "extra_headers": strict.get("extra_headers") or [],
        "order_mismatch": bool(strict.get("order_mismatch")),
        "message": strict.get("message") or "",
        "missing_required_keys": [],
        "unmapped_headers": strict.get("extra_headers") or [],
        "inherited": schema.get("effective_from") == "inherited",
    }


def preview_excel(
    user_id: str, project_id: str, suite_id: str, file_obj: BinaryIO
) -> dict[str, Any]:
    headers, body = _read_excel_table(file_obj)
    return preview_from_headers(
        user_id, project_id, suite_id, headers, sample_rows=body
    )


def preview_workbench(
    user_id: str,
    project_id: str,
    suite_id: str,
    *,
    lanhu_pid: str,
    lanhu_doc_id: str,
    lanhu_page_id: str,
) -> dict[str, Any]:
    headers, body, src = _load_workbench_table(
        user_id,
        lanhu_pid=lanhu_pid,
        lanhu_doc_id=lanhu_doc_id,
        lanhu_page_id=lanhu_page_id,
    )
    result = preview_from_headers(
        user_id, project_id, suite_id, headers, sample_rows=body
    )
    result["page_name"] = src.get("page_name") or ""
    return result


def _map_row_to_case(
    *,
    headers: list[str],
    row: list[Any],
    columns: list[dict[str, Any]],
    mapping: dict[str, str],
    suite_id: str,
    source: str,
    source_ref: str | None = None,
) -> Optional[dict[str, Any]]:
    by_header = schema_db.row_values_by_header(headers, row)
    key_to_role = {str(c.get("key")): str(c.get("role") or "custom") for c in columns}
    fields: dict[str, str] = {}
    for src_label, key in mapping.items():
        fields[key] = by_header.get(src_label, "")

    title = ""
    for key, role in key_to_role.items():
        if role == "title":
            title = str(fields.get(key) or "").strip()
            break
    if not title:
        return None

    priority = "P2"
    status = "draft" if source == "excel" else "ready"
    precondition = ""
    tags: list[str] = []
    step_text = ""
    expect_text = ""

    for key, role in key_to_role.items():
        val = str(fields.get(key) or "").strip()
        if role == "priority":
            p = val.upper() or "P2"
            priority = p if p in ("P0", "P1", "P2", "P3") else "P2"
        elif role == "status":
            s = val.lower() or status
            status = s if s in ("draft", "ready", "deprecated") else status
        elif role == "precondition":
            precondition = val
        elif role == "tags":
            tags = [t.strip() for t in val.split(",") if t.strip()]
        elif role == "module":
            if val and val not in tags:
                tags.append(val)
        elif role == "steps":
            step_text = val
        elif role == "expect":
            expect_text = val

    # 保持导入原文：步骤描述 / 预期结果各自整段存储，不做编号拆分与成对合并
    steps = [{"step": step_text, "expect": expect_text}] if (step_text or expect_text) else []

    return {
        "title": title,
        "suite_id": suite_id,
        "priority": priority,
        "status": status,
        "precondition": precondition,
        "steps": steps,
        "tags": tags,
        "fields": fields,
        "source": source,
        "source_ref": source_ref,
    }


def _resolve_mapping_and_schema(
    user_id: str,
    project_id: str,
    suite_id: str,
    headers: list[str],
    *,
    column_mapping: Optional[dict[str, Any]],
    confirm_schema: bool,
    proposed_columns: Optional[list[dict[str, Any]]],
    source_kind: str,
    source_ref: str,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    schema = schema_db.ensure_inherited_schema(user_id, suite_id)
    locked = str(schema.get("status") or "") == "locked" and bool(schema.get("columns"))

    if not locked:
        cols = proposed_columns or schema_db.build_columns_from_headers(headers)
        schema_db.assert_schema_column_limit(cols)
        if not confirm_schema:
            raise ValueError("请先确认并锁定本目录列结构")
        schema_db.lock_schema(
            user_id,
            project_id,
            suite_id,
            columns=cols,
            source_kind=source_kind,
            source_ref=source_ref,
        )
        # 首导：同名映射
        mapping = {h: str(c.get("key")) for c, h in zip(cols, headers) if str(h).strip()}
        # 更稳：按 label 匹配
        mapping = {}
        for h in headers:
            key = schema_db.match_header_to_column(h, cols)
            if key:
                mapping[h] = key
        mapping = schema_db.validate_mapping(cols, mapping)
        return cols, mapping

    cols = list(schema.get("columns") or [])
    # 已锁定：只允许表头与锁定列完全一致，不再接受字段映射改写
    strict = schema_db.compare_locked_headers_strict(headers, cols)
    if strict.get("mode") != "align":
        raise ValueError(
            str(strict.get("message") or "源表头与已锁定表头不一致，请修改后重新导入")
        )
    mapping = schema_db.validate_mapping(cols, strict.get("suggested_mapping") or {})
    return cols, mapping


def import_excel_with_schema(
    user_id: str,
    project_id: str,
    file_obj: BinaryIO,
    *,
    suite_id: str,
    column_mapping: Optional[dict[str, Any]] = None,
    confirm_schema: bool = False,
    proposed_columns: Optional[list[dict[str, Any]]] = None,
    duplicate_mode: str = "create",
) -> dict[str, Any]:
    target = assert_suite_in_project(user_id, project_id, suite_id)
    # 需要可 seek 的副本以便 preview 已读过时仍可用
    raw = file_obj.read()
    bio = io.BytesIO(raw)
    headers, body = _read_excel_table(bio)
    cols, mapping = _resolve_mapping_and_schema(
        user_id,
        project_id,
        suite_id,
        headers,
        column_mapping=column_mapping,
        confirm_schema=confirm_schema,
        proposed_columns=proposed_columns,
        source_kind="excel",
        source_ref="excel",
    )
    prepared = []
    for row in body:
        item = _map_row_to_case(
            headers=headers,
            row=row,
            columns=cols,
            mapping=mapping,
            suite_id=str(target["id"]),
            source="excel",
        )
        if item:
            prepared.append(item)
    result = _apply_prepared_cases(
        user_id,
        project_id,
        str(target["id"]),
        prepared,
        duplicate_mode=duplicate_mode,
    )
    result["parsed"] = len(prepared)
    result["suite_id"] = str(target["id"])
    result["suite_name"] = target.get("name") or ""
    result["schema_locked"] = True
    return result


def import_workbench_with_schema(
    user_id: str,
    project_id: str,
    *,
    lanhu_pid: str,
    lanhu_doc_id: str,
    lanhu_page_id: str,
    suite_id: str,
    column_mapping: Optional[dict[str, Any]] = None,
    confirm_schema: bool = False,
    proposed_columns: Optional[list[dict[str, Any]]] = None,
    duplicate_mode: str = "create",
) -> dict[str, Any]:
    target = assert_suite_in_project(user_id, project_id, suite_id)
    headers, body, src = _load_workbench_table(
        user_id,
        lanhu_pid=lanhu_pid,
        lanhu_doc_id=lanhu_doc_id,
        lanhu_page_id=lanhu_page_id,
    )
    source_ref = "%s|%s|%s" % (
        src.get("lanhu_pid") or "",
        src.get("lanhu_doc_id") or "",
        src.get("lanhu_page_id") or "",
    )
    cols, mapping = _resolve_mapping_and_schema(
        user_id,
        project_id,
        suite_id,
        headers,
        column_mapping=column_mapping,
        confirm_schema=confirm_schema,
        proposed_columns=proposed_columns,
        source_kind="workbench",
        source_ref=source_ref[:256],
    )
    prepared = []
    for row in body:
        item = _map_row_to_case(
            headers=headers,
            row=row,
            columns=cols,
            mapping=mapping,
            suite_id=str(target["id"]),
            source="workbench",
            source_ref=source_ref[:128],
        )
        if item:
            prepared.append(item)
    result = _apply_prepared_cases(
        user_id,
        project_id,
        str(target["id"]),
        prepared,
        duplicate_mode=duplicate_mode,
    )
    result["parsed"] = len(prepared)
    result["page_name"] = src.get("page_name") or ""
    result["suite_id"] = str(target["id"])
    result["suite_name"] = target.get("name") or ""
    result["schema_locked"] = True
    return result
