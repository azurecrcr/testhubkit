"""将各系统导出的用例 Excel 解析为 jsMind / XMind 可用的思维导图结构。"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from core.services.test_cases.case_template_service import get_case_template_by_id

# 各模板列名 → 语义字段（按模板 columns 精确匹配，辅以别名）
FIELD_CANDIDATES: Dict[str, List[str]] = {
    "name": ["用例名称", "用例标题", "用例摘要", "标题"],
    "module": ["所属模块", "所属产品", "组件", "目录", "模块"],
    "steps": ["步骤描述", "步骤", "测试步骤", "操作步骤"],
    "expected": ["预期结果", "预期"],
    "priority": ["用例等级", "优先级", "用例类型"],
}

MODULE_SEP_RE = re.compile(r"\s*(?:/|>|→|—|–|\\|\|)\s*")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def _normalize_header(value: Any) -> str:
    return _cell_str(value).replace("\n", " ")


def _truncate(text: str, max_len: int = 120) -> str:
    s = _cell_str(text)
    if len(s) <= max_len:
        return s
    return s[: max_len - 1] + "…"


def _split_module_path(module_str: str) -> List[str]:
    s = _cell_str(module_str)
    if not s:
        return ["未分类"]
    parts = [p.strip() for p in MODULE_SEP_RE.split(s) if p.strip()]
    return parts or ["未分类"]


def _find_header_row(ws, template_columns: List[str], max_scan: int = 25) -> int:
    tmpl = {_normalize_header(c) for c in template_columns if c}
    best_row, best_score = 1, 0
    last_row = min(max_scan, ws.max_row or 1)
    for r in range(1, last_row + 1):
        score = 0
        for c in range(1, (ws.max_column or 1) + 1):
            h = _normalize_header(ws.cell(r, c).value)
            if h and h in tmpl:
                score += 1
        if score > best_score:
            best_score = score
            best_row = r
    return best_row if best_score >= 1 else 1


def _read_header_row(ws, row_idx: int) -> List[str]:
    headers: List[str] = []
    for c in range(1, (ws.max_column or 1) + 1):
        headers.append(_normalize_header(ws.cell(row_idx, c).value))
    while headers and not headers[-1]:
        headers.pop()
    return headers


def _match_column_index(headers: List[str], col_name: str) -> Optional[int]:
    target = _normalize_header(col_name)
    if not target:
        return None
    for i, h in enumerate(headers):
        if h == target:
            return i
    for i, h in enumerate(headers):
        if not h:
            continue
        if target in h or h in target:
            return i
    return None


def resolve_field_column_map(
    template_columns: List[str], headers: List[str]
) -> Tuple[Dict[str, int], Dict[str, str]]:
    """返回 { field -> excel列索引(0-based) }, { field -> 匹配到的列名 }"""
    field_idx: Dict[str, int] = {}
    field_label: Dict[str, str] = {}
    used: set[int] = set()

    for field, candidates in FIELD_CANDIDATES.items():
        for cand in candidates:
            if cand not in template_columns:
                continue
            idx = _match_column_index(headers, cand)
            if idx is None or idx in used:
                continue
            field_idx[field] = idx
            field_label[field] = cand
            used.add(idx)
            break

    # 名称列兜底：取模板中第一个能匹配到表头的列
    if "name" not in field_idx:
        for col in template_columns:
            idx = _match_column_index(headers, col)
            if idx is not None and idx not in used:
                field_idx["name"] = idx
                field_label["name"] = col
                used.add(idx)
                break

    return field_idx, field_label


def _parse_rows_from_sheet(ws, header_row: int, field_idx: Dict[str, int]) -> List[Dict[str, str]]:
    headers = _read_header_row(ws, header_row)
    rows: List[Dict[str, str]] = []
    name_idx = field_idx.get("name")
    if name_idx is None:
        return rows

    for r in range(header_row + 1, (ws.max_row or 0) + 1):
        cells = [_cell_str(ws.cell(r, c + 1).value) for c in range(len(headers))]
        if not any(cells):
            continue
        name = cells[name_idx] if name_idx < len(cells) else ""
        if not name:
            continue
        row: Dict[str, str] = {"name": name}
        for field, idx in field_idx.items():
            if field == "name":
                continue
            row[field] = cells[idx] if idx < len(cells) else ""
        rows.append(row)
    return rows


def _stable_branch_id(path: str) -> str:
    s = str(path or "")
    h = 0
    for ch in s:
        h = ((h << 5) - h + ord(ch)) & 0xFFFFFFFF
    return f"ctm_br_{abs(h) % 10**9}"


def _ensure_branch(parent: Dict[str, Any], key: str, full_path: str) -> Dict[str, Any]:
    if key not in parent:
        parent[key] = {"_order": [], "_path": full_path, "_id": _stable_branch_id(full_path)}
        parent["_order"].append(key)
    return parent[key]


def _build_branch_children(branch_map: Dict[str, Any]) -> List[Dict[str, Any]]:
    nodes: List[Dict[str, Any]] = []
    for key in branch_map.get("_order", []):
        entry = branch_map.get(key)
        if entry and isinstance(entry, dict):
            nodes.append(_build_branch_node(key, entry))
    for leaf in branch_map.get("_leaves", []):
        nodes.append(leaf)
    return nodes


def _build_branch_node(title: str, entry: Dict[str, Any]) -> Dict[str, Any]:
    child_nodes: List[Dict[str, Any]] = []
    for ck in entry.get("_order", []):
        sub = entry.get(ck)
        if not sub:
            continue
        if sub.get("_leaf"):
            child_nodes.append(sub["_leaf"])
        else:
            child_nodes.append(_build_branch_node(ck, sub))
    for leaf in entry.get("_leaves", []):
        child_nodes.append(leaf)
    return {
        "id": entry.get("_id") or _stable_branch_id(entry.get("_path", title)),
        "topic": title,
        "expanded": True,
        "data": {"tcType": "module", "moduleName": entry.get("_path", title)},
        "children": child_nodes,
    }


def _make_leaf_node(row: Dict[str, str], row_index: int, module_path: str) -> Dict[str, Any]:
    name = _truncate(row.get("name", ""), 64)
    priority = _cell_str(row.get("priority", ""))
    topic = f"{name} [{priority}]" if priority else name
    children: List[Dict[str, Any]] = []
    steps = _cell_str(row.get("steps", ""))
    expected = _cell_str(row.get("expected", ""))
    if steps:
        children.append(
            {
                "id": f"ctm_step_{row_index}",
                "topic": "步骤：" + _truncate(steps, 200),
                "expanded": True,
                "children": [],
            }
        )
    if expected:
        children.append(
            {
                "id": f"ctm_exp_{row_index}",
                "topic": "预期：" + _truncate(expected, 200),
                "expanded": True,
                "children": [],
            }
        )
    return {
        "id": f"ctm_leaf_{row_index}",
        "topic": topic,
        "expanded": True,
        "data": {"tcType": "leaf", "moduleName": module_path, "rowIndex": row_index},
        "children": children,
    }


def _table_rows_to_semantic_rows(
    columns: List[str],
    rows: List[List[Any]],
) -> Tuple[List[Dict[str, str]], Dict[str, str]]:
    if not columns:
        raise ValueError("缺少表头列定义")

    headers = [_normalize_header(c) for c in columns]
    field_idx, field_labels = resolve_field_column_map(columns, headers)

    if "name" not in field_idx:
        used = set(field_idx.values())
        for field, candidates in FIELD_CANDIDATES.items():
            if field in field_idx:
                continue
            for cand in candidates:
                idx = _match_column_index(headers, cand)
                if idx is not None and idx not in used:
                    field_idx[field] = idx
                    field_labels[field] = headers[idx] or cand
                    used.add(idx)
                    break

    if "name" not in field_idx:
        raise ValueError("未能识别用例名称列，请确认表头包含「用例名称」或同类列名")

    parsed: List[Dict[str, str]] = []
    col_count = len(columns)
    for row in rows:
        if not isinstance(row, (list, tuple)):
            continue
        cells = [_cell_str(row[i]) if i < len(row) else "" for i in range(col_count)]
        name_idx = field_idx["name"]
        if not cells[name_idx]:
            continue
        item: Dict[str, str] = {"name": cells[name_idx]}
        for field, idx in field_idx.items():
            if field == "name":
                continue
            item[field] = cells[idx] if idx < len(cells) else ""
        parsed.append(item)
    return parsed, field_labels


def convert_table_to_mindmap(
    columns: List[str],
    rows: List[List[Any]],
    root_topic: str = "测试用例",
) -> Dict[str, Any]:
    """将工作台表格列+行转为 jsMind 思维导图结构。"""
    parsed_rows, field_labels = _table_rows_to_semantic_rows(columns, rows)
    if not parsed_rows:
        raise ValueError("未解析到有效用例行（需包含用例名称）")

    title = _cell_str(root_topic) or "测试用例"
    if not title.endswith(f"（{len(parsed_rows)}）"):
        title = f"{title}（{len(parsed_rows)}）"

    mind = build_mindmap_from_rows(parsed_rows, root_topic=title)
    return {
        "field_map": field_labels,
        "stats": {
            "total": len(parsed_rows),
            "with_module": sum(1 for r in parsed_rows if _cell_str(r.get("module"))),
            "with_steps": sum(1 for r in parsed_rows if _cell_str(r.get("steps"))),
            "with_expected": sum(1 for r in parsed_rows if _cell_str(r.get("expected"))),
        },
        "mind": mind,
    }


def build_mindmap_from_rows(rows: List[Dict[str, str]], root_topic: str = "测试用例") -> Dict[str, Any]:
    tree_root: Dict[str, Any] = {"_order": [], "_path": ""}
    for i, row in enumerate(rows):
        parts = _split_module_path(row.get("module", ""))
        cursor = tree_root
        path_acc: List[str] = []
        for part in parts:
            path_acc.append(part)
            full_path = " / ".join(path_acc)
            cursor = _ensure_branch(cursor, part, full_path)
        module_path = " / ".join(parts)
        leaf = _make_leaf_node(row, i, module_path)
        if "_leaves" not in cursor:
            cursor["_leaves"] = []
        cursor["_leaves"].append(leaf)

    children = _build_branch_children(tree_root)
    title = _cell_str(root_topic) or "测试用例"
    return {
        "meta": {"name": "TestHub", "author": "TestHub", "version": "1.0"},
        "format": "node_tree",
        "data": {
            "id": "ctm_root",
            "topic": title,
            "expanded": True,
            "data": {"tcType": "root"},
            "children": children,
        },
    }


def convert_excel_to_mindmap(
    file_bytes: bytes,
    filename: str,
    template_id: str,
) -> Dict[str, Any]:
    template = get_case_template_by_id(template_id)
    if not template:
        raise ValueError("未找到所选模板，请刷新页面后重试")

    template_columns: List[str] = list(template.get("columns") or [])
    if not template_columns:
        raise ValueError("模板列配置为空")

    if not filename.lower().endswith((".xlsx", ".xls")):
        raise ValueError("请上传 Excel 文件（.xlsx 或 .xls）")

    if filename.lower().endswith(".xls"):
        raise ValueError("暂不支持旧版 .xls，请另存为 .xlsx 后上传")

    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{exc}") from exc

    try:
        ws = wb.active
        if not ws:
            raise ValueError("Excel 中没有可用的工作表")

        header_row = _find_header_row(ws, template_columns)
        headers = _read_header_row(ws, header_row)
        field_idx, field_labels = resolve_field_column_map(template_columns, headers)

        if "name" not in field_idx:
            raise ValueError(
                "未能识别用例名称列。请确认表头与所选模板一致，且首行或前几行包含列名。"
            )

        rows = _parse_rows_from_sheet(ws, header_row, field_idx)
        if not rows:
            raise ValueError("未解析到有效用例行（需包含用例名称）")

        root_topic = f"{template.get('name') or '测试用例'}（{len(rows)}）"
        mind = build_mindmap_from_rows(rows, root_topic=root_topic)

        return {
            "template": template,
            "field_map": field_labels,
            "header_row": header_row,
            "stats": {
                "total": len(rows),
                "with_module": sum(1 for r in rows if _cell_str(r.get("module"))),
                "with_steps": sum(1 for r in rows if _cell_str(r.get("steps"))),
                "with_expected": sum(1 for r in rows if _cell_str(r.get("expected"))),
            },
            "mind": mind,
        }
    finally:
        wb.close()
