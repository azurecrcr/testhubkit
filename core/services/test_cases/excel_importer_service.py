import ast

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def parse_test_cases(test_cases_text: str):
    test_cases_str = test_cases_text
    if "test_cases" in test_cases_str:
        test_cases_start = test_cases_str.find("[")
        test_cases_end = test_cases_str.rfind("]") + 1
        if test_cases_start == -1 or test_cases_end == -1:
            raise ValueError("无法解析测试用例数据格式，请确保是Python列表格式")
        test_cases_list_str = test_cases_str[test_cases_start:test_cases_end]
        return ast.literal_eval(test_cases_list_str)
    return ast.literal_eval(test_cases_str)


def fill_test_cases_to_excel(filepath: str, test_cases: list):
    def is_row_empty(sheet, row_idx):
        for col_idx in range(1, sheet.max_column + 1):
            if sheet.cell(row=row_idx, column=col_idx).value not in [None, ""]:
                return False
        return True

    wb = load_workbook(filepath)
    ws = wb.active
    wrap_style = Alignment(wrap_text=True)
    current_max_row = ws.max_row
    case_index = 0
    row_idx = 3
    filled_count = 0

    while case_index < len(test_cases):
        if row_idx > current_max_row or is_row_empty(ws, row_idx):
            for col_idx, value in enumerate(test_cases[case_index], start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.value = value.replace("<br>", "\n") if isinstance(value, str) else value
                cell.alignment = wrap_style
            case_index += 1
            filled_count += 1
        row_idx += 1

    if ws.max_column >= 8:
        ws.column_dimensions["E"].width = 25
        ws.column_dimensions["F"].width = 35
        ws.column_dimensions["H"].width = 35
        ws.column_dimensions["G"].width = 20

    wb.save(filepath)
    return filled_count
